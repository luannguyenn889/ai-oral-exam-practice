import glob
import hashlib
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
import warnings
from typing import Any, Optional

import chromadb
import docx
from sentence_transformers import SentenceTransformer

from app.config import settings

warnings.filterwarnings("ignore", message=".*unauthenticated requests.*")
warnings.filterwarnings("ignore", message=".*HF Hub.*")
os.environ["HF_HUB_DISABLE_SYMLINKS_WARNING"] = "1"
os.environ["TOKENIZERS_PARALLELISM"] = "false"

GEMINI_MIN_REQUEST_INTERVAL_SECONDS = 0.7
GEMINI_EMBED_MAX_RETRIES = 3

_embedding_model: Optional[SentenceTransformer] = None
_reranker_model: Any = None
_chroma_client: Optional[chromadb.ClientAPI] = None
_collections: dict[str, Any] = {}
_bm25_cache: dict[str, tuple[Any, list[dict[str, Any]]]] = {}

_gemini_client: Any = None
_gemini_client_api_key: Optional[str] = None
_gemini_request_lock = threading.Lock()
_gemini_last_request_at = 0.0

FILLER_PATTERNS = [
    r"\bdạ\b",
    r"\bthưa thầy\b",
    r"\bthưa cô\b",
    r"\bem xin phép\b",
    r"\btheo em nghĩ\b",
    r"\btheo em nhớ\b",
    r"\bem nghĩ là\b",
    r"\bừm\b",
    r"\bà\b",
    r"\bthì là mà\b",
    r"\bạ\b",
]


def normalize_course_id(course_id: Optional[str]) -> str:
    cleaned = "".join(c for c in (course_id or "default").strip() if c.isalnum() or c in ("-", "_"))
    return cleaned or "default"


def get_active_embedding_name(provider: Optional[str] = None) -> str:
    prov = (provider or settings.embedding_provider or "local").strip().lower()
    if prov == "gemini":
        return settings.gemini_embedding_model_name
    return settings.embedding_model_name


def clean_oral_query(text: str) -> str:
    """Lọc bỏ các từ đệm, ngập ngừng hội thoại để làm sắc nét véc-tơ truy vấn."""
    if not text:
        return ""
    cleaned = text
    for pat in FILLER_PATTERNS:
        cleaned = re.sub(pat, " ", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned if len(cleaned) > 5 else text.strip()


def tokenize_text(text: str) -> list[str]:
    """Tách từ đơn giản cho BM25 (hỗ trợ cả tiếng Việt và thuật ngữ CNTT)."""
    if not text:
        return []
    return re.findall(r"\b\w+\b", text.lower())


def get_embedding_model() -> SentenceTransformer:
    """Tải embedding model theo cơ chế Lazy load từ cache máy tính."""
    global _embedding_model
    if _embedding_model is None:
        model_name = settings.embedding_model_name
        try:
            _embedding_model = SentenceTransformer(model_name, local_files_only=True)
        except Exception:
            _embedding_model = SentenceTransformer(model_name)
    return _embedding_model


def _gemini_embed_content(contents: list[Any]) -> Any:
    """Call Gemini embeddings with gentle pacing and bounded 429 retries."""
    global _gemini_last_request_at
    from google.genai import types

    for attempt in range(GEMINI_EMBED_MAX_RETRIES + 1):
        try:
            with _gemini_request_lock:
                elapsed = time.monotonic() - _gemini_last_request_at
                if elapsed < GEMINI_MIN_REQUEST_INTERVAL_SECONDS:
                    time.sleep(GEMINI_MIN_REQUEST_INTERVAL_SECONDS - elapsed)
                _gemini_last_request_at = time.monotonic()
                return _gemini_client.models.embed_content(
                    model=settings.gemini_embedding_model_name,
                    contents=contents,
                    config=types.EmbedContentConfig(
                        output_dimensionality=settings.gemini_embedding_dimensions
                    ),
                )
        except Exception as exc:
            status_code = getattr(exc, "code", None) or getattr(exc, "status_code", None)
            message = str(exc)
            is_rate_limit = (
                status_code == 429 or "429" in message or "RESOURCE_EXHAUSTED" in message
            )
            if not is_rate_limit or attempt >= GEMINI_EMBED_MAX_RETRIES:
                raise

            retry_match = re.search(
                r"retryDelay['\"]?\s*:\s*['\"]?(\d+(?:\.\d+)?)\s*(ms|s)", message, re.I
            )
            if not retry_match:
                retry_match = re.search(
                    r"retry in\s+(\d+(?:\.\d+)?)\s*(?:seconds?|s)\b", message, re.I
                )
            if retry_match:
                delay = float(retry_match.group(1))
                if len(retry_match.groups()) > 1 and retry_match.group(2).lower() == "ms":
                    delay /= 1000
            else:
                delay = min(5 * (2**attempt), 30)
            delay = min(max(delay, 1.0), 120.0)
            time.sleep(delay)


def embed_texts(
    texts: list[str], task: str = "document", provider: Optional[str] = None
) -> list[list[float]]:
    """Create normalized vectors locally or through Google Gemini embeddings."""
    if not texts:
        return []
    prov = (provider or settings.embedding_provider or "local").strip().lower()
    if prov == "local":
        model = get_embedding_model()
        return model.encode(
            texts, batch_size=32, show_progress_bar=False, normalize_embeddings=True
        ).tolist()

    api_key = (settings.gemini_api_key or os.environ.get("GEMINI_API_KEY", "")).strip()
    if not api_key:
        raise ValueError("Chưa cấu hình GEMINI_API_KEY trong .env để dùng Gemini Embedding.")

    global _gemini_client, _gemini_client_api_key
    if _gemini_client is None or _gemini_client_api_key != api_key:
        from google import genai

        _gemini_client = genai.Client(api_key=api_key)
        _gemini_client_api_key = api_key

    from google.genai import types

    prefix = "task: search result | query: " if task == "query" else "title: none | text: "
    vectors: list[list[float]] = []
    for start in range(0, len(texts), 32):
        contents = [
            types.Content(parts=[types.Part.from_text(text=f"{prefix}{text}")])
            for text in texts[start : start + 32]
        ]
        response = _gemini_embed_content(contents)
        vectors.extend(embedding.values for embedding in (response.embeddings or []))
    if len(vectors) != len(texts):
        raise RuntimeError(f"Gemini trả về {len(vectors)} vector cho {len(texts)} đoạn văn bản.")
    return vectors


def get_reranker_model() -> Any:
    """Tải Cross-Encoder Reranker model để tái xếp hạng Top-K ứng viên."""
    global _reranker_model
    if _reranker_model is None:
        try:
            from sentence_transformers import CrossEncoder

            try:
                _reranker_model = CrossEncoder(
                    settings.reranker_model_name, local_files_only=True
                )
            except Exception:
                _reranker_model = CrossEncoder(settings.reranker_model_name)
        except Exception:
            _reranker_model = None
    return _reranker_model


def get_collection_name(provider: Optional[str] = None) -> str:
    prov = (provider or settings.embedding_provider or "local").strip().lower()
    if prov == "gemini":
        return f"{settings.collection_prefix}_gemini"
    return settings.collection_prefix


def get_chroma_collection(provider: Optional[str] = None) -> Any:
    """Khởi tạo ChromaDB client và collection."""
    global _chroma_client
    db_dir = str(settings.resolved_chroma_db_dir)
    os.makedirs(db_dir, exist_ok=True)
    if _chroma_client is None:
        _chroma_client = chromadb.PersistentClient(path=db_dir)
    coll_name = get_collection_name(provider)
    if coll_name not in _collections:
        _collections[coll_name] = _chroma_client.get_or_create_collection(
            name=coll_name,
            metadata={"hnsw:space": "cosine"},
        )
    return _collections[coll_name]


def invalidate_bm25_cache(course_id: Optional[str] = None) -> None:
    if course_id:
        norm = normalize_course_id(course_id)
        _bm25_cache.pop(norm, None)
    else:
        _bm25_cache.clear()


def get_bm25_index(
    course_id: str = "default", provider: Optional[str] = None
) -> tuple[Any, list[dict[str, Any]]]:
    """Khởi tạo hoặc lấy chỉ mục BM25 của một học phần (course_id) từ ChromaDB."""
    norm_course = normalize_course_id(course_id)
    if norm_course in _bm25_cache:
        return _bm25_cache[norm_course]

    collection = get_chroma_collection(provider)
    try:
        course_data = collection.get(
            where={"course_id": norm_course},
            include=["documents", "metadatas"],
        )
    except Exception:
        course_data = {"ids": [], "documents": [], "metadatas": []}

    bm25_docs: list[dict[str, Any]] = []
    if course_data and course_data.get("documents"):
        for chunk_id, doc, meta in zip(
            course_data["ids"], course_data["documents"], course_data["metadatas"]
        ):
            meta = meta or {}
            bm25_docs.append(
                {
                    "text": doc,
                    "id": chunk_id,
                    "source": meta.get("source", "Không rõ"),
                    "course_id": meta.get("course_id", norm_course),
                    "index": meta.get("index", 0),
                }
            )

    try:
        from rank_bm25 import BM25Okapi

        indexed_docs = []
        tokenized_corpus = []
        for item in bm25_docs:
            tokens = tokenize_text(item["text"])
            if tokens:
                indexed_docs.append(item)
                tokenized_corpus.append(tokens)
        bm25_index = BM25Okapi(tokenized_corpus) if tokenized_corpus else None
        _bm25_cache[norm_course] = (bm25_index, indexed_docs)
        return bm25_index, indexed_docs
    except Exception:
        return None, []


def extract_text_from_docx(file_path: str) -> str:
    """Trích xuất toàn bộ văn bản và bảng biểu từ file .docx."""
    try:
        doc = docx.Document(file_path)
        content_parts: list[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                content_parts.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                unique_cells: list[str] = []
                for c in row_cells:
                    if not unique_cells or c != unique_cells[-1]:
                        unique_cells.append(c)
                if unique_cells:
                    content_parts.append(" | ".join(unique_cells))

        return "\n".join(content_parts)
    except Exception:
        return ""


def extract_text_from_txt(file_path: str) -> str:
    """Trích xuất văn bản từ file text thuần."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    except Exception:
        return ""


def extract_text_from_pdf(file_path: str) -> str:
    """Đọc PDF có text layer (PyMuPDF -> pypdf); nếu là PDF scan thì OCR qua Tesseract."""
    fitz_doc = None
    try:
        import fitz

        fitz_doc = fitz.open(file_path)
        pages_text = []
        for page in fitz_doc:
            t = page.get_text().strip()
            if t:
                pages_text.append(t)
        if pages_text:
            fitz_doc.close()
            return "\n\n".join(pages_text)
    except Exception:
        if fitz_doc:
            fitz_doc.close()

    if fitz_doc:
        fitz_doc.close()

    try:
        import pypdf

        reader = pypdf.PdfReader(file_path)
        pages_text = []
        for page in reader.pages:
            t = page.extract_text()
            if t and t.strip():
                pages_text.append(t.strip())
        if pages_text:
            return "\n\n".join(pages_text)
    except Exception:
        pass

    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError(
            "PDF này không có lớp văn bản đọc được. Cần cài PyMuPDF và Tesseract OCR."
        ) from exc

    tesseract_cmd = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")
    if not tesseract_cmd:
        local_app_data = os.environ.get("LOCALAPPDATA", "")
        for candidate in (
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
            os.path.join(local_app_data, "Programs", "Tesseract-OCR", "tesseract.exe"),
        ):
            if os.path.isfile(candidate):
                tesseract_cmd = candidate
                break
    if not tesseract_cmd:
        raise RuntimeError(
            "PDF này là bản scan không có lớp văn bản và máy chưa cài Tesseract OCR."
        )

    language_result = subprocess.run(
        [tesseract_cmd, "--list-langs"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=True,
    )
    available_languages = set(language_result.stdout.splitlines()[1:])
    languages = [lang for lang in ("vie", "eng") if lang in available_languages]
    if not languages:
        raise RuntimeError("Tesseract chưa có dữ liệu OCR tiếng Việt hoặc tiếng Anh.")
    language = "+".join(languages)
    doc = fitz.open(file_path)
    try:
        pages_text = []
        matrix = fitz.Matrix(2, 2)
        with tempfile.TemporaryDirectory(prefix="rag_pdf_ocr_") as temp_dir:
            for page_number, page in enumerate(doc, 1):
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                image_path = os.path.join(temp_dir, f"page_{page_number:05d}.png")
                pix.save(image_path)
                result = subprocess.run(
                    [
                        tesseract_cmd,
                        image_path,
                        "stdout",
                        "-l",
                        language,
                        "--oem",
                        "1",
                        "--psm",
                        "3",
                    ],
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    check=True,
                )
                page_text = result.stdout.strip()
                if page_text:
                    pages_text.append(page_text)
    finally:
        doc.close()
    if pages_text:
        return "\n\n".join(pages_text)
    raise RuntimeError("Tesseract đã chạy nhưng không nhận diện được chữ trong PDF.")


def get_clean_overlap(text: str, max_chars: int = 100) -> str:
    """Lấy phần đuôi để overlap nhưng bắt buộc giữ trọn vẹn từ ngữ (không cắt ngang từ)."""
    if not text or len(text) <= max_chars:
        return text
    tail = text[-max_chars:].strip()
    first_space = tail.find(" ")
    if first_space != -1 and first_space < len(tail) - 1:
        cleaned = tail[first_space + 1 :].strip()
        cleaned = cleaned.lstrip(".,:;|-)]} ")
        return cleaned
    return tail


def chunk_text(
    text: str,
    source_name: str,
    course_id: str = "default",
    chunk_size: int = 600,
    overlap: int = 100,
) -> list[dict[str, Any]]:
    """
    Chia văn bản thành các đoạn nhỏ (chunks) theo ranh giới đoạn văn và câu hoàn chỉnh.
    Đảm bảo không bị cắt cụt từ ngữ ở đầu hoặc cuối mỗi đoạn trích xuất.
    """
    if not text:
        return []

    norm_course = normalize_course_id(course_id)
    paragraphs = text.split("\n")
    chunks: list[dict[str, Any]] = []
    current_chunk = ""
    chunk_index = 1

    for p in paragraphs:
        p = p.strip()
        if not p:
            continue

        if len(current_chunk) + len(p) + 1 <= chunk_size:
            current_chunk = f"{current_chunk}\n{p}" if current_chunk else p
        else:
            if current_chunk:
                clean_entry = current_chunk.strip().lstrip(".,:;|-)]} ")
                if len(clean_entry) > 30:
                    chunks.append(
                        {
                            "id": f"{norm_course}::{source_name}_chunk_{chunk_index}",
                            "text": clean_entry,
                            "source": source_name,
                            "course_id": norm_course,
                            "index": chunk_index,
                        }
                    )
                    chunk_index += 1

                tail_overlap = get_clean_overlap(current_chunk, max_chars=overlap)
                current_chunk = f"{tail_overlap}\n{p}" if tail_overlap else p
            else:
                sentences = re.split(r"(?<=[.?!])\s+", p)
                sub_chunk = ""
                for s in sentences:
                    if len(sub_chunk) + len(s) + 1 <= chunk_size:
                        sub_chunk = f"{sub_chunk} {s}".strip()
                    else:
                        if sub_chunk:
                            clean_s = sub_chunk.strip().lstrip(".,:;|-)]} ")
                            chunks.append(
                                {
                                    "id": f"{norm_course}::{source_name}_chunk_{chunk_index}",
                                    "text": clean_s,
                                    "source": source_name,
                                    "course_id": norm_course,
                                    "index": chunk_index,
                                }
                            )
                            chunk_index += 1
                        sub_chunk = s
                current_chunk = sub_chunk

    if current_chunk.strip():
        clean_entry = current_chunk.strip().lstrip(".,:;|-)]} ")
        if len(clean_entry) > 20:
            chunks.append(
                {
                    "id": f"{norm_course}::{source_name}_chunk_{chunk_index}",
                    "text": clean_entry,
                    "source": source_name,
                    "course_id": norm_course,
                    "index": chunk_index,
                }
            )

    return chunks


def list_course_files(course_id: str = "default") -> list[str]:
    course_dir = str(settings.course_documents_dir(course_id))
    all_files: list[str] = []
    for pattern in ("**/*.docx", "**/*.txt", "**/*.pdf"):
        for path in glob.glob(os.path.join(course_dir, pattern), recursive=True):
            if not os.path.basename(path).startswith("~$"):
                all_files.append(path)
    return sorted(set(all_files))


def index_documents(
    course_id: str = "default",
    force_reindex: bool = False,
    embedding_provider: Optional[str] = None,
) -> dict[str, Any]:
    """
    Đồng bộ tài liệu trong data/documents/{course_id}/ với ChromaDB.
    Chỉ quét và cập nhật các chunk thuộc đúng course_id.
    """
    norm_course = normalize_course_id(course_id)
    course_dir = str(settings.course_documents_dir(norm_course))
    collection = get_chroma_collection(embedding_provider)
    active_embedding_name = get_active_embedding_name(embedding_provider)

    if force_reindex:
        invalidate_bm25_cache(norm_course)

    all_files = list_course_files(norm_course)

    try:
        existing = collection.get(
            where={"course_id": norm_course},
            include=["metadatas"],
        )
    except Exception:
        existing = {"ids": [], "metadatas": []}

    existing_by_source: dict[str, dict[str, Any]] = {}
    for chunk_id, meta in zip(existing.get("ids", []), existing.get("metadatas", [])):
        meta = meta or {}
        source = meta.get("source")
        if source:
            item = existing_by_source.setdefault(
                source,
                {
                    "ids": [],
                    "content_hash": meta.get("content_hash"),
                    "embedding_model": meta.get("embedding_model"),
                },
            )
            item["ids"].append(chunk_id)

    current_sources: set[str] = set()
    added_or_updated = 0
    removed = 0

    for file_path in all_files:
        source = os.path.relpath(file_path, course_dir).replace("\\", "/")
        current_sources.add(source)
        try:
            with open(file_path, "rb") as source_file:
                content_hash = hashlib.sha256(source_file.read()).hexdigest()
        except OSError:
            continue

        previous = existing_by_source.get(source)
        if (
            not force_reindex
            and previous
            and previous["content_hash"] == content_hash
            and previous.get("embedding_model") == active_embedding_name
        ):
            continue

        extension = os.path.splitext(file_path)[1].lower()
        if extension == ".docx":
            raw_text = extract_text_from_docx(file_path)
        elif extension == ".pdf":
            raw_text = extract_text_from_pdf(file_path)
        else:
            raw_text = extract_text_from_txt(file_path)

        file_chunks = chunk_text(raw_text, source_name=source, course_id=norm_course)
        if file_chunks:
            old_ids = previous["ids"] if previous else []
            texts = [chunk["text"] for chunk in file_chunks]
            embeddings = embed_texts(texts, task="document", provider=embedding_provider)
            collection.upsert(
                ids=[chunk["id"] for chunk in file_chunks],
                documents=texts,
                metadatas=[
                    {
                        "source": source,
                        "course_id": norm_course,
                        "index": chunk["index"],
                        "content_hash": content_hash,
                        "embedding_model": active_embedding_name,
                    }
                    for chunk in file_chunks
                ],
                embeddings=embeddings,
            )
            stale_ids = set(old_ids) - {chunk["id"] for chunk in file_chunks}
            if stale_ids:
                collection.delete(ids=list(stale_ids))
            added_or_updated += len(file_chunks)

    for source, item in existing_by_source.items():
        if source not in current_sources and item["ids"]:
            collection.delete(ids=item["ids"])
            removed += len(item["ids"])

    invalidate_bm25_cache(norm_course)

    try:
        after_course = collection.get(where={"course_id": norm_course}, include=[])
        course_chunk_count = len(after_course.get("ids", []))
    except Exception:
        course_chunk_count = 0

    return {
        "course_id": norm_course,
        "status": "ready" if course_chunk_count > 0 else "empty",
        "total_files": len(all_files),
        "added_or_updated_chunks": added_or_updated,
        "removed_chunks": removed,
        "total_chunks": course_chunk_count,
        "embedding_model": active_embedding_name,
    }


def query_relevant_chunks(
    query_text: str,
    course_id: str = "default",
    top_k: int = 3,
    score_threshold: Optional[float] = None,
    search_mode: str = "hybrid",
    use_rerank: bool = False,
    embedding_provider: Optional[str] = None,
) -> list[dict[str, Any]]:
    """
    Truy vấn nâng cao các đoạn giáo trình liên quan nhất (Hybrid Dense + BM25 + RRF + Optional Reranker)
    được cô lập theo từng học phần (course_id).
    """
    if not query_text or not query_text.strip():
        return []
    if search_mode not in {"dense", "bm25", "hybrid"}:
        raise ValueError("search_mode phải là 'dense', 'bm25' hoặc 'hybrid'.")
    if top_k < 1:
        return []

    norm_course = normalize_course_id(course_id)
    collection = get_chroma_collection(embedding_provider)

    try:
        existing_course = collection.get(where={"course_id": norm_course}, include=[])
        course_count = len(existing_course.get("ids", []))
    except Exception:
        course_count = 0

    if course_count == 0:
        index_documents(
            course_id=norm_course,
            force_reindex=False,
            embedding_provider=embedding_provider,
        )
        try:
            existing_course = collection.get(where={"course_id": norm_course}, include=[])
            course_count = len(existing_course.get("ids", []))
        except Exception:
            course_count = 0

    if course_count == 0:
        return []

    cleaned_query = clean_oral_query(query_text)
    candidate_pool_size = min(max(top_k * 3, 8), course_count)
    candidates: dict[str, dict[str, Any]] = {}

    # 1. TÌM KIẾM BẰNG DENSE VECTOR (CHROMADB)
    if search_mode in ("hybrid", "dense"):
        query_vector = embed_texts([cleaned_query], task="query", provider=embedding_provider)
        results = collection.query(
            query_embeddings=query_vector,
            n_results=candidate_pool_size,
            where={"course_id": norm_course},
            include=["documents", "metadatas", "distances"],
        )
        if results and results.get("documents"):
            docs = results["documents"][0]
            metas = results["metadatas"][0] if results.get("metadatas") else [{}] * len(docs)
            distances = (
                results["distances"][0] if results.get("distances") else [0.0] * len(docs)
            )
            ids = results.get("ids", [[]])[0]
            for rank, (chunk_id, doc, meta, dist) in enumerate(
                zip(ids, docs, metas, distances), 1
            ):
                meta = meta or {}
                sim = round(1.0 - float(dist), 4)
                candidates[chunk_id] = {
                    "id": chunk_id,
                    "text": doc,
                    "source": meta.get("source", "Không rõ"),
                    "course_id": meta.get("course_id", norm_course),
                    "index": meta.get("index", 0),
                    "dense_sim": sim,
                    "dense_distance": round(float(dist), 4),
                    "dense_rank": rank,
                    "bm25_score": None,
                    "bm25_rank": 999,
                    "rrf_score": 0.0,
                }

    # 2. TÌM KIẾM BẰNG SPARSE BM25 (TỪ KHÓA VIẾT TẮT / CHUYÊN NGÀNH)
    if search_mode in ("hybrid", "bm25"):
        bm25, bm25_docs = get_bm25_index(norm_course, provider=embedding_provider)
        if bm25 and bm25_docs:
            tokenized_q = tokenize_text(cleaned_query)
            if tokenized_q:
                bm25_scores = bm25.get_scores(tokenized_q)
                top_bm25_indices = sorted(
                    range(len(bm25_scores)), key=lambda i: bm25_scores[i], reverse=True
                )[:candidate_pool_size]

                for rank, idx in enumerate(top_bm25_indices, 1):
                    score = float(bm25_scores[idx])
                    if score <= 0:
                        continue
                    b_doc = bm25_docs[idx]
                    chunk_id = b_doc["id"]
                    if chunk_id not in candidates:
                        candidates[chunk_id] = {
                            "id": chunk_id,
                            "text": b_doc["text"],
                            "source": b_doc["source"],
                            "course_id": b_doc.get("course_id", norm_course),
                            "index": b_doc["index"],
                            "dense_sim": None,
                            "dense_rank": 999,
                            "bm25_score": round(score, 2),
                            "bm25_rank": rank,
                            "rrf_score": 0.0,
                        }
                    else:
                        candidates[chunk_id]["bm25_score"] = round(score, 2)
                        candidates[chunk_id]["bm25_rank"] = rank

    if not candidates:
        return []

    # 3. KẾT HỢP ĐIỂM SỐ BẰNG THUẬT TOÁN RRF (RECIPROCAL RANK FUSION, k=60)
    candidate_list = list(candidates.values())
    k_constant = 60.0
    for c in candidate_list:
        retrieval_sources = []
        if c.get("dense_rank", 999) != 999:
            retrieval_sources.append("dense")
        if c.get("bm25_rank", 999) != 999:
            retrieval_sources.append("BM25")
        c["retrieved_by"] = " + ".join(retrieval_sources) or "unknown"
        score_dense = (
            1.0 / (k_constant + c["dense_rank"]) if c.get("dense_rank", 999) < 999 else 0.0
        )
        score_bm25 = (
            1.0 / (k_constant + c["bm25_rank"]) if c.get("bm25_rank", 999) < 999 else 0.0
        )
        c["rrf_score"] = round(score_dense + score_bm25, 5)

    candidate_list.sort(key=lambda x: x["rrf_score"], reverse=True)

    # 4. GIAI ĐOẠN 2: TÁI XẾP HẠNG (CROSS-ENCODER RE-RANKING)
    if use_rerank and len(candidate_list) > 1:
        reranker = get_reranker_model()
        if reranker is not None:
            pairs = [[cleaned_query, c["text"]] for c in candidate_list]
            try:
                rerank_scores = reranker.predict(pairs)
                for c, r_score in zip(candidate_list, rerank_scores):
                    c["rerank_score"] = round(float(r_score), 4)
                    c["method"] = f"{search_mode.title()} + Cross-Encoder Rerank"
                candidate_list.sort(key=lambda x: x["rerank_score"], reverse=True)
            except Exception:
                for c in candidate_list:
                    c["method"] = f"{search_mode.title()} RRF"
        else:
            for c in candidate_list:
                c["method"] = f"{search_mode.title()} RRF"
    else:
        for c in candidate_list:
            c["method"] = f"{search_mode.title()} RRF"

    if score_threshold is not None:
        candidate_list = [
            c
            for c in candidate_list
            if c.get("dense_rank", 999) == 999
            or (c.get("dense_sim") is not None and c["dense_sim"] >= score_threshold)
        ]

    return candidate_list[:top_k]


def format_rag_context(chunks: list[dict[str, Any]]) -> str:
    """Định dạng danh sách các đoạn trích xuất thành văn bản sạch sẽ để chèn vào Prompt của LLM."""
    if not chunks:
        return ""
    formatted = []
    for i, c in enumerate(chunks, 1):
        method_str = f" | Phương thức: {c.get('method', 'RAG')}" if "method" in c else ""
        sim_str = (
            f" (Cosine: {c.get('dense_sim', 'N/A')}; RRF: {c.get('rrf_score', 'N/A')}; "
            f"Rerank: {c.get('rerank_score', 'N/A')}{method_str})"
        )
        formatted.append(
            f"[Đoạn trích {i} - ID: {c.get('id', 'N/A')} - Nguồn: {c['source']} - "
            f"Chunk: {c.get('index', 'N/A')}{sim_str}]\n{c['text']}"
        )
    return "\n\n---\n\n".join(formatted)


def get_rag_status(
    course_id: str = "default", embedding_provider: Optional[str] = None
) -> dict[str, Any]:
    """Lấy thông tin trạng thái kho tri thức RAG của một học phần."""
    norm_course = normalize_course_id(course_id)
    course_dir = str(settings.course_documents_dir(norm_course))
    try:
        collection = get_chroma_collection(embedding_provider)
        existing = collection.get(where={"course_id": norm_course}, include=[])
        count = len(existing.get("ids", []))

        all_paths = list_course_files(norm_course)
        files = [os.path.relpath(p, course_dir).replace("\\", "/") for p in all_paths]

        bm25, _ = get_bm25_index(norm_course, provider=embedding_provider) if count else (None, [])
        has_bm25 = bm25 is not None

        return {
            "course_id": norm_course,
            "status": "ready" if count > 0 else "empty",
            "total_chunks": count,
            "total_files": len(files),
            "files": files,
            "embedding_model": get_active_embedding_name(embedding_provider),
            "reranker_model": settings.reranker_model_name,
            "hybrid_bm25_active": has_bm25,
            "reranker_enabled_by_default": False,
            "search_modes_available": ["dense"] + (["bm25", "hybrid"] if has_bm25 else []),
            "features": [
                "Boundary-Preserved Chunking (600 chars, 100 overlap)",
                "Course-level Isolation (course_id metadata filter)",
                "ChromaDB HNSW Dense Cosine Search",
                *(["BM25 Okapi Sparse Keyword Search", "Reciprocal Rank Fusion (RRF)"] if has_bm25 else []),
                "Cross-Encoder available (opt-in for experiment)",
            ],
            "error": None,
        }
    except Exception as e:
        return {
            "course_id": norm_course,
            "status": "error",
            "total_chunks": 0,
            "total_files": 0,
            "files": [],
            "embedding_model": get_active_embedding_name(embedding_provider),
            "reranker_model": settings.reranker_model_name,
            "hybrid_bm25_active": False,
            "reranker_enabled_by_default": False,
            "search_modes_available": [],
            "features": [],
            "error": str(e),
        }

